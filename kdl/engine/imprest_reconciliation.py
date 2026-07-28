"""Standalone Imprest reconciliation workbook generator.

This engine intentionally has no dependency on either Imprest loader engine.
It translates the supplied VBA cleaning and one-to-one matching workflow into
an auditable, non-destructive Python/openpyxl process.
"""

from __future__ import annotations

from collections import defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import os
from pathlib import Path
import re
import tempfile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from kdl.tabular_import import load_workbook_from_source


CANONICAL_HEADERS = (
    "Source",
    "Category",
    "GL Date",
    "Event Class",
    "Transaction Number",
    "Line Description",
    "Debit",
    "Credit",
)

HEADER_ALIASES = {
    "Source": {"source", "journalsource"},
    "Category": {"category", "journalcategory"},
    "GL Date": {"gldate", "accountingdate", "effectivedate", "date"},
    "Event Class": {"eventclass", "eventtypeclass"},
    "Transaction Number": {
        "transactionnumber",
        "transactionno",
        "transactionnum",
        "documentnumber",
    },
    "Line Description": {
        "linedescription",
        "line",
        "description",
        "transactiondescription",
    },
    "Debit": {
        "debit",
        "debitamount",
        "entereddr",
        "accounteddr",
        "dr",
    },
    "Credit": {
        "credit",
        "creditamount",
        "enteredcr",
        "accountedcr",
        "cr",
    },
}

KNOWN_REPORT_LABELS = {
    "account",
    "beginningbalanceforperiod",
    "endofreport",
    "endingbalanceforperiod",
    "gokledger",
    "ledgername",
    "source",
    "subledgeraccounting",
}

TITLE_FILL = PatternFill("solid", fgColor="0B5EA8")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")
MATCH_FILL = PatternFill("solid", fgColor="C6EFCE")
UNMATCHED_FILL = PatternFill("solid", fgColor="FFC7CE")
REVIEW_FILL = PatternFill("solid", fgColor="FFEB9C")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_BLUE = Side(style="thin", color="9EBCD2")
SECTION_BORDER = Border(bottom=Side(style="medium", color="0B5EA8"))
AMOUNT_FORMAT = '#,##0.00;[Red]-#,##0.00'
ID_PATTERN = re.compile(r"\d{5,}")
IMPREST_ID_PATTERN = re.compile(r"\bIMP\s*[-/]?\s*(\d{5,})\b", re.IGNORECASE)


@dataclass
class ImprestReconciliationResult:
    output_path: str
    source_sheet: str
    transaction_count: int
    matched_pairs: int
    unmatched_count: int
    review_count: int
    unmatched_debit: Decimal
    unmatched_credit: Decimal
    difference: Decimal
    warnings: tuple[str, ...]
    message: str


@dataclass
class _Transaction:
    values: list[Any]
    source_row: int
    identifier: str
    side: str
    amount: Decimal | None
    status: str
    pair_id: str = ""
    note: str = ""


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _display_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal_amount(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError):
            return None

    text = str(value).strip()
    if not text or text in {"-", "—"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = (
        text.replace(",", "")
        .replace("KES", "")
        .replace("Ksh", "")
        .replace("KSH", "")
        .strip()
    )
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    if negative:
        amount = -amount
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _extract_identifier(line_description: Any, transaction_number: Any) -> str:
    for value in (line_description, transaction_number):
        text = _display_text(value)
        if not text:
            continue
        imprest_matches = IMPREST_ID_PATTERN.findall(text)
        if imprest_matches:
            return imprest_matches[-1]
        numeric_matches = ID_PATTERN.findall(text)
        if numeric_matches:
            return numeric_matches[-1]
    return ""


def _header_mapping_for_row(values: list[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalised = [_normalise_header(value) for value in values]
    for canonical, aliases in HEADER_ALIASES.items():
        for index, candidate in enumerate(normalised):
            if candidate in aliases:
                mapping[canonical] = index
                break
    return mapping


def _looks_like_data_row(values: list[Any]) -> bool:
    if len(values) < 8:
        return False
    has_reference = bool(_display_text(values[4]) or _display_text(values[5]))
    has_amount = _decimal_amount(values[6]) is not None or _decimal_amount(values[7]) is not None
    return has_reference and has_amount


def _unique_headers(values: list[Any], mapping: dict[str, int], last_col: int) -> list[str]:
    mapped_by_index = {column: name for name, column in mapping.items()}
    headers: list[str] = []
    used: set[str] = set()
    for index in range(last_col):
        base = mapped_by_index.get(index) or _display_text(values[index] if index < len(values) else "")
        base = base or f"Column {index + 1}"
        candidate = base
        suffix = 2
        while candidate.casefold() in used:
            candidate = f"{base} {suffix}"
            suffix += 1
        headers.append(candidate)
        used.add(candidate.casefold())
    return headers


def _locate_data(ws) -> tuple[int, int, dict[str, int], list[str]]:
    scan_rows = min(ws.max_row, 250)
    scan_cols = min(max(ws.max_column, 8), 100)
    best: tuple[int, dict[str, int], list[Any]] | None = None

    for row_number in range(1, scan_rows + 1):
        values = [ws.cell(row_number, column).value for column in range(1, scan_cols + 1)]
        mapping = _header_mapping_for_row(values)
        required = {"Debit", "Credit"}
        has_reference = "Line Description" in mapping or "Transaction Number" in mapping
        if required.issubset(mapping) and has_reference and len(mapping) >= 5:
            best = (row_number, mapping, values)
            break

    if best is not None:
        header_row, mapping, header_values = best
        last_header_col = max(
            [index + 1 for index, value in enumerate(header_values) if _display_text(value)]
            + [max(mapping.values()) + 1]
        )
        last_col = max(last_header_col, min(ws.max_column, 250))
        headers = _unique_headers(header_values, mapping, last_col)
        return header_row + 1, last_col, mapping, headers

    # Compatibility fallback for the original macro's fixed A:H report layout.
    for row_number in range(1, scan_rows + 1):
        values = [ws.cell(row_number, column).value for column in range(1, 9)]
        if _looks_like_data_row(values):
            mapping = {name: index for index, name in enumerate(CANONICAL_HEADERS)}
            last_col = max(8, min(ws.max_column, 250))
            headers = list(CANONICAL_HEADERS) + [
                f"Column {index}" for index in range(9, last_col + 1)
            ]
            return row_number, last_col, mapping, headers

    raise ValueError(
        "Could not find the transaction columns. The source must contain Debit, "
        "Credit and either Line Description or Transaction Number columns, or use "
        "the original eight-column A:H layout."
    )


def _cell_values(ws, row_number: int, last_col: int) -> list[Any]:
    return [ws.cell(row_number, column).value for column in range(1, last_col + 1)]


def _read_transactions(ws) -> tuple[list[_Transaction], list[str], dict[str, int], list[str]]:
    data_start, last_col, mapping, headers = _locate_data(ws)
    transactions: list[_Transaction] = []
    warnings: list[str] = []
    debit_col = mapping["Debit"]
    credit_col = mapping["Credit"]
    description_col = mapping.get("Line Description")
    transaction_col = mapping.get("Transaction Number")

    for row_number in range(data_start, ws.max_row + 1):
        values = _cell_values(ws, row_number, last_col)
        if not any(_display_text(value) for value in values):
            continue

        first_label = _normalise_header(values[0] if values else "")
        debit = _decimal_amount(values[debit_col])
        credit = _decimal_amount(values[credit_col])
        description = values[description_col] if description_col is not None else ""
        transaction_number = values[transaction_col] if transaction_col is not None else ""
        reference_text = _display_text(description) or _display_text(transaction_number)

        if (
            first_label in KNOWN_REPORT_LABELS
            and debit is None
            and credit is None
            and not reference_text
        ):
            continue

        has_other_transaction_data = any(
            _display_text(values[mapping[name]])
            for name in ("Source", "Category", "GL Date", "Event Class")
            if name in mapping
        )
        if debit is None and credit is None and not reference_text and not has_other_transaction_data:
            continue

        identifier = _extract_identifier(description, transaction_number)
        debit_nonzero = debit is not None and debit != 0
        credit_nonzero = credit is not None and credit != 0
        if debit_nonzero:
            values[debit_col] = abs(debit)
        if credit_nonzero:
            values[credit_col] = -abs(credit)
        status = "UNMATCHED"
        side = ""
        amount: Decimal | None = None
        note = ""

        if debit_nonzero and credit_nonzero:
            status = "REVIEW"
            note = "Both Debit and Credit contain non-zero amounts."
        elif debit_nonzero:
            side = "DEBIT"
            amount = abs(debit)
        elif credit_nonzero:
            side = "CREDIT"
            amount = abs(credit)
        else:
            status = "REVIEW"
            note = "No usable Debit or Credit amount."

        if not identifier:
            status = "REVIEW"
            note = (note + " " if note else "") + (
                "No identifier containing five or more digits was found."
            )

        transactions.append(
            _Transaction(
                values=values,
                source_row=row_number,
                identifier=identifier,
                side=side,
                amount=amount,
                status=status,
                note=note,
            )
        )

    if not transactions:
        raise ValueError("No transaction rows were found on the selected sheet.")

    _match_transactions(transactions)
    transactions.sort(
        key=lambda item: (
            item.identifier == "",
            item.identifier.casefold(),
            item.amount if item.amount is not None else Decimal("Infinity"),
            item.source_row,
        )
    )
    review_count = sum(item.status == "REVIEW" for item in transactions)
    if review_count:
        warnings.append(
            f"{review_count} row(s) require review because the identifier or amount was invalid."
        )
    return transactions, headers, mapping, warnings


def _match_transactions(transactions: list[_Transaction]) -> None:
    debit_queues: dict[tuple[str, Decimal], deque[int]] = defaultdict(deque)
    credit_queues: dict[tuple[str, Decimal], deque[int]] = defaultdict(deque)

    for index, item in enumerate(transactions):
        if item.status == "REVIEW" or not item.identifier or item.amount is None:
            continue
        key = (item.identifier.casefold(), item.amount)
        if item.side == "DEBIT":
            debit_queues[key].append(index)
        elif item.side == "CREDIT":
            credit_queues[key].append(index)

    pair_number = 0
    for key in sorted(set(debit_queues) | set(credit_queues)):
        debits = debit_queues[key]
        credits = credit_queues[key]
        while debits and credits:
            pair_number += 1
            pair_id = f"M{pair_number:06d}"
            debit_item = transactions[debits.popleft()]
            credit_item = transactions[credits.popleft()]
            debit_item.status = "MATCHED"
            credit_item.status = "MATCHED"
            debit_item.pair_id = pair_id
            credit_item.pair_id = pair_id
            debit_item.note = "Exact identifier and amount match."
            credit_item.note = "Exact identifier and amount match."

    for item in transactions:
        if item.status == "UNMATCHED":
            item.note = "No opposite-side row has the same identifier and amount."


def _serialisable_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, int, float, str, bool)) or value is None:
        return value
    return str(value)


def _row_output(item: _Transaction) -> list[Any]:
    return [
        *[_serialisable_value(value) for value in item.values],
        item.identifier,
        item.side,
        float(item.amount) if item.amount is not None else None,
        item.status,
        item.pair_id,
        item.note,
        item.source_row,
    ]


def _style_title(ws, title: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.fill = TITLE_FILL
    cell.font = Font(color="FFFFFF", bold=True, size=15)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.sheet_view.showGridLines = False


def _style_summary_label(cell) -> None:
    cell.fill = LABEL_FILL
    cell.font = Font(bold=True, color="1F4E78")
    cell.border = SECTION_BORDER


def _style_header_row(ws, row_number: int, last_col: int) -> None:
    for cell in ws.iter_cols(
        min_col=1, max_col=last_col, min_row=row_number, max_row=row_number
    ):
        header = cell[0]
        header.fill = HEADER_FILL
        header.font = WHITE_FONT
        header.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.border = Border(bottom=THIN_BLUE)
    ws.row_dimensions[row_number].height = 32


def _fit_columns(ws, headers: list[str], rows: list[list[Any]]) -> None:
    for column_number, header in enumerate(headers, 1):
        sample_values = [_display_text(row[column_number - 1]) for row in rows[:250]]
        width = max([len(str(header))] + [len(value) for value in sample_values]) + 2
        lowered = header.casefold()
        if "description" in lowered or "note" in lowered:
            width = min(max(width, 28), 52)
        elif "path" in lowered:
            width = min(max(width, 28), 45)
        else:
            width = min(max(width, 11), 24)
        ws.column_dimensions[get_column_letter(column_number)].width = width


def _add_table(ws, start_row: int, end_row: int, last_col: int, name: str) -> None:
    if end_row <= start_row:
        return
    ref = f"A{start_row}:{ws.cell(end_row, last_col).coordinate}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _apply_data_formats(
    ws,
    first_data_row: int,
    last_data_row: int,
    mapping: dict[str, int],
    output_headers: list[str],
) -> None:
    if last_data_row < first_data_row:
        return
    for canonical in ("Debit", "Credit"):
        column = mapping.get(canonical)
        if column is not None:
            cells = next(ws.iter_cols(
                min_col=column + 1,
                max_col=column + 1,
                min_row=first_data_row,
                max_row=last_data_row,
            ))
            for cell in cells:
                cell.number_format = AMOUNT_FORMAT
    if "GL Date" in mapping:
        column = mapping["GL Date"] + 1
        cells = next(ws.iter_cols(
            min_col=column,
            max_col=column,
            min_row=first_data_row,
            max_row=last_data_row,
        ))
        for cell in cells:
            cell.number_format = "dd-mmm-yyyy"

    match_amount_col = output_headers.index("Match Amount") + 1
    cells = next(ws.iter_cols(
        min_col=match_amount_col,
        max_col=match_amount_col,
        min_row=first_data_row,
        max_row=last_data_row,
    ))
    for cell in cells:
        cell.number_format = AMOUNT_FORMAT


def _create_main_sheet(
    wb: Workbook,
    transactions: list[_Transaction],
    source_headers: list[str],
    mapping: dict[str, int],
    source_path: str,
    source_sheet: str,
) -> tuple[list[str], int]:
    ws = wb.active
    ws.title = "Imprest_Reconciliation"
    output_headers = [
        *source_headers,
        "Reconciliation ID",
        "Side",
        "Match Amount",
        "Match Status",
        "Match Pair",
        "Review Note",
        "Source Row",
    ]
    last_col = len(output_headers)
    _style_title(ws, "Imprest Reconciliation", last_col)

    summary = (
        ("Source File", os.path.basename(source_path), "Transactions", len(transactions)),
        ("Source Sheet", source_sheet, "Matched Pairs", sum(i.status == "MATCHED" for i in transactions) // 2),
        ("Matching Rule", "Final identifier + exact amount + opposite side", "Unmatched", sum(i.status == "UNMATCHED" for i in transactions)),
        ("Source Protection", "Original workbook was not changed", "Review Items", sum(i.status == "REVIEW" for i in transactions)),
    )
    for row_offset, (label1, value1, label2, value2) in enumerate(summary, 3):
        ws.cell(row_offset, 1, label1)
        ws.cell(row_offset, 2, value1)
        ws.merge_cells(
            start_row=row_offset,
            start_column=2,
            end_row=row_offset,
            end_column=5,
        )
        ws.cell(row_offset, 6, label2)
        ws.cell(row_offset, 7, value2)
        _style_summary_label(ws.cell(row_offset, 1))
        _style_summary_label(ws.cell(row_offset, 6))

    header_row = 8
    for column, header in enumerate(output_headers, 1):
        ws.cell(header_row, column, header)
    _style_header_row(ws, header_row, last_col)

    data_rows = [_row_output(item) for item in transactions]
    for row_number, (item, row_values) in enumerate(
        zip(transactions, data_rows), header_row + 1
    ):
        for column, value in enumerate(row_values, 1):
            ws.cell(row_number, column, value)
        if item.status == "MATCHED":
            for cell in ws[row_number][:last_col]:
                cell.fill = MATCH_FILL

    last_data_row = header_row + len(transactions)
    _add_table(ws, header_row, last_data_row, last_col, "ImprestReconciliationTable")
    _apply_data_formats(ws, header_row + 1, last_data_row, mapping, output_headers)
    _fit_columns(ws, output_headers, data_rows)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 20)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:{ws.cell(last_data_row, last_col).coordinate}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    return output_headers, header_row


def _create_unmatched_sheet(
    wb: Workbook,
    transactions: list[_Transaction],
    output_headers: list[str],
    mapping: dict[str, int],
) -> None:
    ws = wb.create_sheet("Unmatched_Report")
    last_col = len(output_headers)
    _style_title(ws, "Unmatched and Review Transactions", last_col)
    ws["A3"] = "Unmatched"
    ws["B3"] = sum(item.status == "UNMATCHED" for item in transactions)
    ws["D3"] = "Review Items"
    ws["E3"] = sum(item.status == "REVIEW" for item in transactions)
    _style_summary_label(ws["A3"])
    _style_summary_label(ws["D3"])

    header_row = 6
    for column, header in enumerate(output_headers, 1):
        ws.cell(header_row, column, header)
    _style_header_row(ws, header_row, last_col)

    exceptions = [item for item in transactions if item.status != "MATCHED"]
    data_rows = [_row_output(item) for item in exceptions]
    for row_number, (item, row_values) in enumerate(
        zip(exceptions, data_rows), header_row + 1
    ):
        for column, value in enumerate(row_values, 1):
            ws.cell(row_number, column, value)
        fill = REVIEW_FILL if item.status == "REVIEW" else UNMATCHED_FILL
        for cell in ws[row_number][:last_col]:
            cell.fill = fill

    last_data_row = header_row + len(exceptions)
    if exceptions:
        _add_table(ws, header_row, last_data_row, last_col, "ImprestExceptionsTable")
        _apply_data_formats(ws, header_row + 1, last_data_row, mapping, output_headers)

        debit_col = mapping["Debit"] + 1
        credit_col = mapping["Credit"] + 1
        total_row = last_data_row + 2
        difference_row = total_row + 1
        ws.cell(total_row, max(1, debit_col - 1), "Total Unmatched")
        ws.cell(total_row, debit_col, f"=SUM({ws.cell(header_row + 1, debit_col).coordinate}:{ws.cell(last_data_row, debit_col).coordinate})")
        ws.cell(total_row, credit_col, f"=SUM({ws.cell(header_row + 1, credit_col).coordinate}:{ws.cell(last_data_row, credit_col).coordinate})")
        ws.cell(difference_row, max(1, debit_col - 1), "Difference")
        ws.cell(
            difference_row,
            debit_col,
            f"={ws.cell(total_row, debit_col).coordinate}+{ws.cell(total_row, credit_col).coordinate}",
        )
        for cell in (
            ws.cell(total_row, max(1, debit_col - 1)),
            ws.cell(difference_row, max(1, debit_col - 1)),
        ):
            cell.font = Font(bold=True)
        for cell in (
            ws.cell(total_row, debit_col),
            ws.cell(total_row, credit_col),
            ws.cell(difference_row, debit_col),
        ):
            cell.number_format = AMOUNT_FORMAT
            cell.font = Font(bold=True)
    else:
        ws.cell(header_row + 1, 1, "All valid transactions were matched.")

    _fit_columns(ws, output_headers, data_rows)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _totals(
    transactions: list[_Transaction],
    mapping: dict[str, int],
) -> tuple[Decimal, Decimal, Decimal]:
    unmatched_debit = Decimal("0.00")
    unmatched_credit = Decimal("0.00")
    for item in transactions:
        if item.status == "MATCHED":
            continue
        debit = _decimal_amount(item.values[mapping["Debit"]])
        credit = _decimal_amount(item.values[mapping["Credit"]])
        if debit is not None:
            unmatched_debit += debit
        if credit is not None:
            unmatched_credit += credit
    return unmatched_debit, unmatched_credit, unmatched_debit + unmatched_credit


def suggest_imprest_reconciliation_output(source_path: str) -> str:
    source = Path(source_path)
    return str(source.with_name(f"{source.stem}_Imprest_Reconciliation.xlsx"))


def create_imprest_reconciliation(
    source_path: str,
    output_path: str,
    sheet_name: str,
) -> ImprestReconciliationResult:
    """Create a two-sheet Imprest reconciliation workbook without changing the source."""

    source_path = os.path.abspath(source_path)
    output_path = os.path.abspath(output_path)
    if not os.path.isfile(source_path):
        raise FileNotFoundError(f"Source workbook was not found: {source_path}")
    if os.path.normcase(source_path) == os.path.normcase(output_path):
        raise ValueError("The output path must be different from the source workbook.")

    source_wb = load_workbook_from_source(
        source_path,
        sheet_names=[sheet_name] if sheet_name else None,
        data_only=True,
        read_only=False,
        keep_links=False,
    )
    try:
        if sheet_name not in source_wb.sheetnames:
            raise ValueError(f"Worksheet '{sheet_name}' was not found in the source workbook.")
        transactions, headers, mapping, warnings = _read_transactions(source_wb[sheet_name])
    finally:
        close = getattr(source_wb, "close", None)
        if callable(close):
            close()

    workbook = Workbook()
    output_headers, _ = _create_main_sheet(
        workbook,
        transactions,
        headers,
        mapping,
        source_path,
        sheet_name,
    )
    _create_unmatched_sheet(workbook, transactions, output_headers, mapping)
    workbook.properties.title = "Imprest Reconciliation"
    workbook.properties.subject = "Exact debit/credit matching by final identifier and amount"
    workbook.properties.creator = "NT DL"
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass

    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    temp_handle = tempfile.NamedTemporaryFile(
        prefix=".imprest_reconciliation_",
        suffix=".xlsx",
        dir=output_dir or None,
        delete=False,
    )
    temp_path = temp_handle.name
    temp_handle.close()
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.unlink(temp_path)

    matched_pairs = sum(item.status == "MATCHED" for item in transactions) // 2
    unmatched_count = sum(item.status == "UNMATCHED" for item in transactions)
    review_count = sum(item.status == "REVIEW" for item in transactions)
    unmatched_debit, unmatched_credit, difference = _totals(transactions, mapping)
    message_lines = [
        "Imprest reconciliation completed.",
        f"Source sheet: {sheet_name}",
        f"Transactions analysed: {len(transactions):,}",
        f"Exact matched pairs: {matched_pairs:,}",
        f"Unmatched transactions: {unmatched_count:,}",
        f"Review items: {review_count:,}",
        f"Unmatched debit: {unmatched_debit:,.2f}",
        f"Unmatched credit: {unmatched_credit:,.2f}",
        f"Difference: {difference:,.2f}",
        "",
        "The original workbook was not changed.",
    ]
    if warnings:
        message_lines.extend(("", "Review warnings:", *[f"- {item}" for item in warnings]))

    return ImprestReconciliationResult(
        output_path=output_path,
        source_sheet=sheet_name,
        transaction_count=len(transactions),
        matched_pairs=matched_pairs,
        unmatched_count=unmatched_count,
        review_count=review_count,
        unmatched_debit=unmatched_debit,
        unmatched_credit=unmatched_credit,
        difference=difference,
        warnings=tuple(warnings),
        message="\n".join(message_lines),
    )
