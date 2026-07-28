"""Standalone IFMIS transaction filtering and cleaning engine."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import os
from pathlib import Path
import re
import tempfile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from kdl.tabular_import import load_workbook_from_source, resolve_workbook_sheet_name


OUTPUT_HEADERS = (
    "Source",
    "Category",
    "GL Date",
    "Event Class",
    "Transaction Number",
    "Line Description",
    "Debit",
    "Credit",
    "Count",
)

RECONCILIATION_HEADERS = (
    "Account",
    "Description",
    "Period",
    "Opening Balance",
    "Credit",
    "Payment / Debit",
    "Calculated Closing",
    "IFMIS Closing",
    "Difference",
    "Status",
    "Source Rows",
)

HEADER_ALIASES = {
    "Source": {"source", "journalsource"},
    "Category": {"category", "journalcategory"},
    "GL Date": {"gldate", "accountingdate", "effectivedate", "date"},
    "Event Class": {"eventclass", "eventtypeclass"},
    "Transaction Number": {
        "transactionnumber",
        "transactionno",
        "transactionnum",
        "documentnumber",
    },
    "Line Description": {
        "linedescription",
        "line",
        "description",
        "transactiondescription",
    },
    "Debit": {
        "debit",
        "debitamount",
        "entereddr",
        "accounteddr",
        "dr",
    },
    "Credit": {
        "credit",
        "creditamount",
        "enteredcr",
        "accountedcr",
        "cr",
    },
}

KNOWN_REPORT_LABELS = {
    "",
    "account",
    "beginningbalanceforperiod",
    "endofreport",
    "endingbalanceforperiod",
    "gokledger",
    "ledgername",
    "source",
    "subledgeraccounting",
}

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="0B5EA8")
TOTAL_FILL = PatternFill("solid", fgColor="D9EAF7")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
CHECK_FILL = PatternFill("solid", fgColor="FFC7CE")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_BLUE = Side(style="thin", color="9EBCD2")
AMOUNT_FORMAT = '#,##0.00;[Red]-#,##0.00'
JULY_REVIEW_FONT = Font(color="C00000")


@dataclass
class FilterResult:
    output_path: str
    source_sheet: str
    row_count: int
    debit_total: Decimal
    credit_total: Decimal
    reconciliation_count: int
    reconciliation_pass_count: int
    reconciliation_exception_count: int
    message: str


@dataclass
class _BalanceReconciliation:
    account: str
    description: str
    period: str
    opening_balance: Decimal
    credit: Decimal
    payment: Decimal
    closing_balance: Decimal
    beginning_row: int
    period_total_row: int
    ending_row: int

    @property
    def calculated_closing(self) -> Decimal:
        return self.opening_balance + self.credit + self.payment

    @property
    def difference(self) -> Decimal:
        return self.calculated_closing - self.closing_balance

    @property
    def status(self) -> str:
        return "PASS" if abs(self.difference) <= Decimal("0.01") else "CHECK"


def _normalise_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _display_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decimal_amount(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError):
            return None

    text = str(value).strip()
    if not text or text in {"-", "—"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = (
        text.replace(",", "")
        .replace("KES", "")
        .replace("Ksh", "")
        .replace("KSH", "")
        .strip()
    )
    try:
        amount = Decimal(text)
    except InvalidOperation:
        return None
    if negative:
        amount = -amount
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _header_mapping(values: list[Any]) -> dict[str, int]:
    normalised = [_normalise_header(value) for value in values]
    mapping: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for index, candidate in enumerate(normalised):
            if candidate in aliases:
                mapping[canonical] = index
                break
    return mapping


def _looks_like_data_row(values: list[Any]) -> bool:
    if len(values) < 8:
        return False
    has_transaction = bool(_display_text(values[4]) or _display_text(values[5]))
    has_amount = _decimal_amount(values[6]) is not None or _decimal_amount(values[7]) is not None
    return has_transaction and has_amount


def _locate_data(ws) -> tuple[int, dict[str, int]]:
    scan_rows = min(ws.max_row, 250)
    scan_cols = min(max(ws.max_column, 8), 100)
    for row_number in range(1, scan_rows + 1):
        values = [ws.cell(row_number, column).value for column in range(1, scan_cols + 1)]
        mapping = _header_mapping(values)
        if (
            {"Debit", "Credit"}.issubset(mapping)
            and ("Transaction Number" in mapping or "Line Description" in mapping)
            and len(mapping) >= 5
        ):
            return row_number + 1, mapping

    # Compatibility with the original VBA report whose transaction data begins
    # after a fixed banner block and always occupies A:H.
    for row_number in range(1, scan_rows + 1):
        values = [ws.cell(row_number, column).value for column in range(1, 9)]
        if _looks_like_data_row(values):
            return row_number, {
                name: index for index, name in enumerate(OUTPUT_HEADERS[:8])
            }

    raise ValueError(
        "Could not find the transaction columns. The source must contain Debit, "
        "Credit and either Transaction Number or Line Description columns, or use "
        "the original eight-column A:H layout."
    )


def _serialisable_value(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, int, float, str, bool)) or value is None:
        return value
    return str(value)


def _is_july_transaction(value: Any) -> bool:
    parsed: date | None = None
    if isinstance(value, datetime):
        parsed = value.date()
    elif isinstance(value, date):
        parsed = value
    elif isinstance(value, str):
        text = value.strip()
        for pattern in ("%d-%b-%Y", "%d-%b-%y", "%b %d, %Y", "%Y-%m-%d"):
            try:
                parsed = datetime.strptime(text, pattern).date()
                break
            except ValueError:
                continue
    return bool(parsed and parsed.month == 7)


def _read_filtered_rows(ws) -> tuple[list[list[Any]], Decimal, Decimal]:
    data_start, mapping = _locate_data(ws)
    rows: list[list[Any]] = []
    debit_total = Decimal("0.00")
    credit_total = Decimal("0.00")

    for row_number in range(data_start, ws.max_row + 1):
        source_values = [
            ws.cell(row_number, mapping[header] + 1).value for header in OUTPUT_HEADERS[:8]
        ]
        if not any(_display_text(value) for value in source_values):
            continue
        if _normalise_header(source_values[0]) in KNOWN_REPORT_LABELS:
            continue
        if len(_header_mapping(source_values)) >= 5:
            continue

        has_transaction = bool(
            _display_text(source_values[4]) or _display_text(source_values[5])
        )
        debit = _decimal_amount(source_values[6])
        credit = _decimal_amount(source_values[7])
        has_context = any(_display_text(value) for value in source_values[:4])
        if not has_transaction and debit is None and credit is None and not has_context:
            continue

        if debit is not None:
            debit = -abs(debit) if debit != 0 else Decimal("0.00")
            source_values[6] = float(debit)
            debit_total += debit
        if credit is not None:
            credit = abs(credit) if credit != 0 else Decimal("0.00")
            source_values[7] = float(credit)
            credit_total += credit

        rows.append([_serialisable_value(value) for value in source_values])

    if not rows:
        raise ValueError("No transaction rows were found on the selected worksheet.")
    return rows, debit_total, credit_total


def _signed_balance(debit_value: Any, credit_value: Any) -> Decimal:
    debit = _decimal_amount(debit_value) or Decimal("0.00")
    credit = _decimal_amount(credit_value) or Decimal("0.00")
    return credit - debit


def _read_balance_reconciliations(ws) -> list[_BalanceReconciliation]:
    reconciliations: list[_BalanceReconciliation] = []
    current_account = ""
    current_description = ""
    pending: dict[str, Any] | None = None

    for row_number in range(1, ws.max_row + 1):
        values = [ws.cell(row_number, column).value for column in range(1, 5)]
        first = _normalise_header(values[0])
        second = _normalise_header(values[1])

        if first == "account":
            current_account = _display_text(values[1])
            current_description = _display_text(values[3])
            continue

        if first == "beginningbalanceforperiod":
            pending = {
                "account": current_account,
                "description": current_description,
                "period": _display_text(values[1]),
                "opening_balance": _signed_balance(values[2], values[3]),
                "credit": None,
                "payment": None,
                "beginning_row": row_number,
                "period_total_row": None,
            }
            continue

        if second == "periodtotal" and pending is not None:
            debit = _decimal_amount(values[2]) or Decimal("0.00")
            credit = _decimal_amount(values[3]) or Decimal("0.00")
            pending["payment"] = -abs(debit) if debit != 0 else Decimal("0.00")
            pending["credit"] = abs(credit) if credit != 0 else Decimal("0.00")
            pending["period_total_row"] = row_number
            continue

        if first == "endingbalanceforperiod" and pending is not None:
            if pending["credit"] is not None and pending["payment"] is not None:
                reconciliations.append(
                    _BalanceReconciliation(
                        account=pending["account"],
                        description=pending["description"],
                        period=pending["period"],
                        opening_balance=pending["opening_balance"],
                        credit=pending["credit"],
                        payment=pending["payment"],
                        closing_balance=_signed_balance(values[2], values[3]),
                        beginning_row=pending["beginning_row"],
                        period_total_row=pending["period_total_row"],
                        ending_row=row_number,
                    )
                )
            pending = None

    return reconciliations


def _fit_columns(ws, rows: list[list[Any]]) -> None:
    for column_number, header in enumerate(OUTPUT_HEADERS, 1):
        values = [_display_text(row[column_number - 1]) for row in rows[:250]]
        width = max([len(header)] + [len(value) for value in values]) + 2
        if header == "Line Description":
            width = min(max(width, 28), 52)
        elif header in {"Debit", "Credit"}:
            width = min(max(width, 14), 18)
        elif header == "Count":
            width = 11
        else:
            width = min(max(width, 11), 24)
        ws.column_dimensions[ws.cell(1, column_number).column_letter].width = width


def _build_balance_reconciliation_sheet(
    workbook: Workbook,
    reconciliations: list[_BalanceReconciliation],
    source_path: str,
    source_sheet: str,
) -> None:
    ws = workbook.create_sheet("Balance_Reconciliation")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:K1")
    title = ws["A1"]
    title.value = "IFMIS Balance Reconciliation"
    title.fill = TITLE_FILL
    title.font = Font(color="FFFFFF", bold=True, size=15)
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = TITLE_FILL

    pass_count = sum(item.status == "PASS" for item in reconciliations)
    exception_count = len(reconciliations) - pass_count
    model_status = "PASS" if reconciliations and exception_count == 0 else "CHECK"

    ws.merge_cells("B3:F3")
    ws.merge_cells("B4:F4")
    summary_values = {
        "A3": "Formula",
        "B3": "Opening Balance + Credit + Payment / Debit = Closing Balance",
        "G3": "Periods",
        "H3": len(reconciliations),
        "I3": "Passed",
        "J3": pass_count,
        "A4": "Source",
        "B4": f"{os.path.basename(source_path)} [{source_sheet}]",
        "G4": "Exceptions",
        "H4": exception_count,
        "I4": "Model Status",
        "J4": model_status,
    }
    for coordinate, value in summary_values.items():
        ws[coordinate] = value
    for coordinate in ("A3", "G3", "I3", "A4", "G4", "I4"):
        ws[coordinate].font = Font(bold=True, color="1F4E78")
    ws["J4"].fill = PASS_FILL if model_status == "PASS" else CHECK_FILL
    ws["J4"].font = Font(bold=True, color="006100" if model_status == "PASS" else "9C0006")

    header_row = 6
    for column, header in enumerate(RECONCILIATION_HEADERS, 1):
        cell = ws.cell(header_row, column, header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_BLUE)
    ws.row_dimensions[header_row].height = 34

    for row_number, item in enumerate(reconciliations, header_row + 1):
        values = (
            item.account,
            item.description,
            item.period,
            float(item.opening_balance),
            float(item.credit),
            float(item.payment),
            f"=D{row_number}+E{row_number}+F{row_number}",
            float(item.closing_balance),
            f"=G{row_number}-H{row_number}",
            f'=IF(ABS(I{row_number})<=0.01,"PASS","CHECK")',
            f"{item.beginning_row} / {item.period_total_row} / {item.ending_row}",
        )
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
        for column in range(4, 10):
            ws.cell(row_number, column).number_format = AMOUNT_FORMAT
        status_cell = ws.cell(row_number, 10)
        status_cell.fill = PASS_FILL if item.status == "PASS" else CHECK_FILL
        status_cell.font = Font(
            bold=True,
            color="006100" if item.status == "PASS" else "9C0006",
        )
        status_cell.alignment = Alignment(horizontal="center")

    if reconciliations:
        last_row = header_row + len(reconciliations)
        table = Table(
            displayName="BalanceReconciliationTable",
            ref=f"A{header_row}:K{last_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        last_row = header_row + 1
        ws.merge_cells(start_row=last_row, start_column=1, end_row=last_row, end_column=11)
        ws.cell(
            last_row,
            1,
            "No IFMIS beginning, period-total, and ending balance sections were found.",
        )

    widths = (24, 45, 12, 18, 18, 18, 20, 18, 16, 12, 20)
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(header_row, column).column_letter].width = width
    ws.freeze_panes = "A7"
    ws.auto_filter.ref = f"A{header_row}:K{last_row}" if reconciliations else None
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1


def _build_output_workbook(
    rows: list[list[Any]],
    reconciliations: list[_BalanceReconciliation],
    source_path: str,
    source_sheet: str,
) -> Workbook:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Filtered_Data"
    ws.sheet_view.showGridLines = False

    for column, header in enumerate(OUTPUT_HEADERS, 1):
        cell = ws.cell(1, column, header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_BLUE)
    ws.row_dimensions[1].height = 32

    data_last_row = len(rows) + 1
    for row_number, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
        ws.cell(
            row_number,
            9,
            f'=COUNTIF($E$2:$E${data_last_row},E{row_number})',
        )
        if _is_july_transaction(values[2]):
            for cell in ws[row_number][:9]:
                cell.font = JULY_REVIEW_FONT

    total_row = data_last_row + 1
    ws.cell(total_row, 6, "Total")
    ws.cell(total_row, 7, f"=SUM(G2:G{data_last_row})")
    ws.cell(total_row, 8, f"=SUM(H2:H{data_last_row})")
    for cell in ws[total_row][:9]:
        cell.fill = TOTAL_FILL
        cell.font = Font(bold=True, color="1F4E78")

    for row_number in range(2, total_row + 1):
        ws.cell(row_number, 7).number_format = AMOUNT_FORMAT
        ws.cell(row_number, 8).number_format = AMOUNT_FORMAT
        ws.cell(row_number, 3).number_format = "dd-mmm-yyyy"

    table = Table(
        displayName="FilterDataTable",
        ref=f"A1:I{total_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    _fit_columns(ws, [row + [None] for row in rows])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{total_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1

    _build_balance_reconciliation_sheet(
        workbook,
        reconciliations,
        source_path,
        source_sheet,
    )
    workbook.properties.title = "IFMIS Transaction Filter"
    workbook.properties.subject = (
        f"Filtered transaction data from {os.path.basename(source_path)} [{source_sheet}]"
    )
    workbook.properties.creator = "NT DL"
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except AttributeError:
        pass
    return workbook


def suggest_filter_output(source_path: str) -> str:
    source = Path(source_path)
    return str(source.with_name(f"{source.stem}_Filtered.xlsx"))


def create_filtered_workbook(
    source_path: str,
    output_path: str,
    sheet_name: str,
) -> FilterResult:
    """Create a filtered workbook without modifying the selected source."""

    source_path = os.path.abspath(source_path)
    output_path = os.path.abspath(output_path)
    if not os.path.isfile(source_path):
        raise FileNotFoundError(f"Source workbook was not found: {source_path}")
    if os.path.normcase(source_path) == os.path.normcase(output_path):
        raise ValueError("The output path must be different from the source workbook.")

    source_wb = load_workbook_from_source(
        source_path,
        sheet_names=[sheet_name] if sheet_name else None,
        data_only=True,
        read_only=False,
        keep_links=False,
    )
    try:
        resolved_sheet_name = resolve_workbook_sheet_name(
            list(source_wb.sheetnames),
            sheet_name,
        )
        if resolved_sheet_name is None:
            raise ValueError(f"Worksheet '{sheet_name}' was not found in the source workbook.")
        source_ws = source_wb[resolved_sheet_name]
        rows, debit_total, credit_total = _read_filtered_rows(source_ws)
        reconciliations = _read_balance_reconciliations(source_ws)
    finally:
        close = getattr(source_wb, "close", None)
        if callable(close):
            close()

    workbook = _build_output_workbook(
        rows,
        reconciliations,
        source_path,
        resolved_sheet_name,
    )
    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)
    temp_handle = tempfile.NamedTemporaryFile(
        prefix=".transaction_filter_",
        suffix=".xlsx",
        dir=output_dir or None,
        delete=False,
    )
    temp_path = temp_handle.name
    temp_handle.close()
    try:
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
    finally:
        if os.path.exists(temp_path):
            os.unlink(temp_path)

    message = "\n".join(
        (
            "Filter Engine completed.",
            f"Source sheet: {resolved_sheet_name}",
            f"Transaction rows retained: {len(rows):,}",
            f"Payment / Debit total (negative): {debit_total:,.2f}",
            f"Credit total (positive): {credit_total:,.2f}",
            f"Net movement: {debit_total + credit_total:,.2f}",
            f"Balance periods reconciled: {len(reconciliations):,}",
            f"Balance checks passed: {sum(item.status == 'PASS' for item in reconciliations):,}",
            f"Balance exceptions: {sum(item.status != 'PASS' for item in reconciliations):,}",
            "",
            "Opening + Credit + Payment / Debit is checked against IFMIS Closing.",
            "Count formulas were added using Transaction Number.",
            "The original workbook was not changed.",
        )
    )
    return FilterResult(
        output_path=output_path,
        source_sheet=resolved_sheet_name,
        row_count=len(rows),
        debit_total=debit_total,
        credit_total=credit_total,
        reconciliation_count=len(reconciliations),
        reconciliation_pass_count=sum(item.status == "PASS" for item in reconciliations),
        reconciliation_exception_count=sum(
            item.status != "PASS" for item in reconciliations
        ),
        message=message,
    )
