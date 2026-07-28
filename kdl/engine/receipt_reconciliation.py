"""Receipt-only reconciliation for Kenyan F.O. 30 PDF reports.

The public entry point is :func:`create_receipt_reconciliation`.  It extracts
Sections 2 and 4, validates both printed totals, performs one-to-one matching,
and writes the required four-sheet audit workbook.
"""

from __future__ import annotations

import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


SECTION_HEADERS = {
    1: "PAYMENTS IN CASH BOOK NOT YET RECORDED IN BANK STATEMENT",
    2: "RECEIPTS IN BANK STATEMENT NOT YET RECORDED IN CASH BOOK",
    3: "PAYMENTS IN BANK STATEMENT NOT YET RECORDED IN CASH BOOK",
    4: "RECEIPTS IN CASH BOOK NOT YET RECORDED IN BANK STATEMENT",
}
REQUIRED_SHEETS = [
    "Summary",
    "Matched",
    "Bank Outstanding",
    "Cashbook Outstanding",
]

_SECTION_RE = re.compile(r"^\s*([1-4])\.\s+(.+?)\s*$", re.IGNORECASE)
_DATE_AMOUNT_RE = re.compile(
    r"(?P<date>\d{2}-[A-Z]{3}-\d{2})\s+"
    r"(?P<amount>\(?-?[\d,]+\.\d{2}\)?)\s*$",
    re.IGNORECASE,
)
_TOTAL_RE = re.compile(r"Total\s*:\s*([\d,]+\.\d{2})", re.IGNORECASE)
_SUFFIX_RE = re.compile(r"^[A-Z0-9]{1,3}$", re.IGNORECASE)
_DATE_RANGE_RE = re.compile(
    r"From\s+Date\s*:\s*(\d{2}-[A-Z]{3}-\d{2})\s+"
    r"To\s*:\s*(\d{2}-[A-Z]{3}-\d{2})",
    re.IGNORECASE,
)
_BANK_RE = re.compile(
    r"Bank\s*:\s*(.*?)\s*,\s*Branch\s*:\s*(.*?)\s*,\s*"
    r"Account\s+Number\s*:\s*([A-Z0-9-]+)",
    re.IGNORECASE,
)

_NAVY = "17365D"
_BLUE = "2F75B5"
_PALE_BLUE = "D9EAF7"
_SUBTITLE = "F3F6F9"
_GREEN = "E2F0D9"
_AMBER = "FFF2CC"
_RED = "FCE4D6"
_WHITE = "FFFFFF"
_GRID = Side(style="thin", color="C9D2DC")
_BORDER = Border(left=_GRID, right=_GRID, top=_GRID, bottom=_GRID)
_CURRENCY_FORMAT = '#,##0.00;[Red]-#,##0.00'
_DATE_FORMAT = "dd-mmm-yyyy"
_MONTH_FORMAT = "mmm yyyy"


class ReceiptReconciliationError(RuntimeError):
    """Raised when extraction or quality-control requirements are not met."""


@dataclass
class Receipt:
    section: int
    reference: str
    transaction_date: date
    amount: Decimal
    pdf_row: int
    source_pages: list[int] = field(default_factory=list)

    @property
    def normalized_reference(self) -> str:
        return normalize_reference(self.reference)

    @property
    def month(self) -> date:
        return self.transaction_date.replace(day=1)

    @property
    def source_page(self) -> str:
        pages = sorted(set(self.source_pages))
        if not pages:
            return ""
        if len(pages) == 1:
            return f"Page {pages[0]}"
        if pages == list(range(pages[0], pages[-1] + 1)):
            return f"Pages {pages[0]}-{pages[-1]}"
        return "Pages " + ", ".join(str(page) for page in pages)


@dataclass
class Match:
    bank: Receipt
    cashbook: Receipt
    status: str
    reason: str
    confidence: float
    match_id: str = ""


@dataclass
class ReconciliationResult:
    output_path: str
    bank_count: int
    cashbook_count: int
    exact_count: int
    probable_count: int
    bank_outstanding_count: int
    cashbook_outstanding_count: int
    bank_total: Decimal
    cashbook_total: Decimal

    @property
    def message(self) -> str:
        return "\n".join(
            [
                "Receipt reconciliation created successfully.",
                f"Section 2 receipts: {self.bank_count}  |  {self.bank_total:,.2f}",
                f"Section 4 receipts: {self.cashbook_count}  |  {self.cashbook_total:,.2f}",
                f"Exact matches: {self.exact_count}",
                f"Probable matches: {self.probable_count}",
                f"Bank outstanding: {self.bank_outstanding_count}",
                f"Cash-book outstanding: {self.cashbook_outstanding_count}",
                f"Workbook: {self.output_path}",
            ]
        )


@dataclass
class _Metadata:
    source_name: str
    bank: str
    branch: str
    account_number: str
    period_from: date
    period_to: date
    printed_totals: dict[int, Decimal]


def normalize_reference(reference: str) -> str:
    """Normalize only presentation whitespace and case, never substantive text."""
    return re.sub(r"\s+", "", str(reference or "").upper())


def _is_blank_reference(reference: str) -> bool:
    normalized = normalize_reference(reference)
    return normalized in {
        "",
        "-",
        "N/A",
        "NA",
        "NIL",
        "NONE",
        "BLANK",
        "BNKISBLANK",
        "REFERENCEISBLANK",
    } or normalized.endswith("ISBLANK")


def _parse_date(value: str) -> date:
    try:
        return datetime.strptime(value.upper(), "%d-%b-%y").date()
    except ValueError as exc:
        raise ReceiptReconciliationError(f"Invalid receipt date in PDF: {value}") from exc


def _parse_amount(value: str) -> Decimal:
    clean = value.strip().replace(",", "")
    negative = clean.startswith("(") and clean.endswith(")")
    clean = clean.strip("()")
    try:
        amount = Decimal(clean).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise ReceiptReconciliationError(f"Invalid receipt amount in PDF: {value}") from exc
    return -amount if negative else amount


def _extract_metadata(first_page_text: str, source_name: str) -> _Metadata:
    range_match = _DATE_RANGE_RE.search(first_page_text)
    bank_match = _BANK_RE.search(first_page_text)
    if not range_match:
        raise ReceiptReconciliationError(
            "Could not find the F.O. 30 From/To reconciliation period."
        )
    if not bank_match:
        raise ReceiptReconciliationError(
            "Could not find the bank, branch and account number in the F.O. 30 PDF."
        )

    printed_totals: dict[int, Decimal] = {}
    for section in (2, 4):
        pattern = re.compile(
            rf"{section}\.\s*{re.escape(SECTION_HEADERS[section])}\s+"
            r"([\d,]+\.\d{2})",
            re.IGNORECASE,
        )
        total_match = pattern.search(first_page_text)
        if total_match:
            printed_totals[section] = _parse_amount(total_match.group(1))

    return _Metadata(
        source_name=source_name,
        bank=bank_match.group(1).strip(),
        branch=bank_match.group(2).strip(),
        account_number=bank_match.group(3).strip(),
        period_from=_parse_date(range_match.group(1)),
        period_to=_parse_date(range_match.group(2)),
        printed_totals=printed_totals,
    )


def _extract_pdf(pdf_path: str) -> tuple[_Metadata, list[Receipt], list[Receipt]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ReceiptReconciliationError(
            "Receipt reconciliation requires pypdf. Reinstall TNT DL to include it."
        ) from exc

    try:
        reader = PdfReader(pdf_path)
    except Exception as exc:
        raise ReceiptReconciliationError(f"Could not open the selected PDF: {exc}") from exc

    if not reader.pages:
        raise ReceiptReconciliationError("The selected PDF has no pages.")

    first_page_text = reader.pages[0].extract_text() or ""
    metadata = _extract_metadata(first_page_text, os.path.basename(pdf_path))
    receipts: dict[int, list[Receipt]] = {2: [], 4: []}
    detail_totals: dict[int, Decimal] = {}
    active_section: int | None = None

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            text = page.extract_text() or ""

        for raw_line in text.splitlines():
            line = raw_line.rstrip()
            section_match = _SECTION_RE.match(line)
            if section_match:
                section_number = int(section_match.group(1))
                active_section = section_number if section_number in (2, 4) else None
                continue

            if active_section not in (2, 4):
                continue

            total_match = _TOTAL_RE.search(line)
            if total_match:
                detail_totals[active_section] = _parse_amount(total_match.group(1))
                active_section = None
                continue

            record_match = _DATE_AMOUNT_RE.search(line)
            if record_match:
                reference = line[: record_match.start()].strip()
                if not reference:
                    continue
                receipts[active_section].append(
                    Receipt(
                        section=active_section,
                        reference=reference,
                        transaction_date=_parse_date(record_match.group("date")),
                        amount=_parse_amount(record_match.group("amount")),
                        pdf_row=len(receipts[active_section]) + 1,
                        source_pages=[page_number],
                    )
                )
                continue

            suffix = line.strip()
            if (
                suffix
                and _SUFFIX_RE.fullmatch(suffix)
                and receipts[active_section]
                and suffix.upper() not in {"NO", "DATE"}
            ):
                latest = receipts[active_section][-1]
                latest.reference = f"{latest.reference}{suffix}"
                if page_number not in latest.source_pages:
                    latest.source_pages.append(page_number)

    for section in (2, 4):
        if not receipts[section]:
            raise ReceiptReconciliationError(
                f"No receipts were extracted from Section {section}."
            )
        if section in detail_totals:
            if section in metadata.printed_totals and (
                detail_totals[section] != metadata.printed_totals[section]
            ):
                raise ReceiptReconciliationError(
                    f"Section {section} total differs between the F.O. 30 summary "
                    f"({metadata.printed_totals[section]:,.2f}) and detail "
                    f"({detail_totals[section]:,.2f})."
                )
            metadata.printed_totals[section] = detail_totals[section]
        if section not in metadata.printed_totals:
            raise ReceiptReconciliationError(
                f"Could not find the printed total for Section {section}."
            )

        extracted_total = sum(
            (receipt.amount for receipt in receipts[section]), Decimal("0.00")
        ).quantize(Decimal("0.01"))
        if extracted_total != metadata.printed_totals[section]:
            difference = extracted_total - metadata.printed_totals[section]
            raise ReceiptReconciliationError(
                f"Section {section} extraction does not reconcile to the PDF. "
                f"Extracted {extracted_total:,.2f}; printed "
                f"{metadata.printed_totals[section]:,.2f}; difference "
                f"{difference:,.2f}. No workbook was created."
            )

    return metadata, receipts[2], receipts[4]


def _candidate_evidence(bank: Receipt, cash: Receipt):
    if bank.amount != cash.amount:
        return None

    bank_ref = bank.normalized_reference
    cash_ref = cash.normalized_reference
    same_reference = bank_ref == cash_ref and bool(bank_ref)
    same_date = bank.transaction_date == cash.transaction_date
    date_difference = abs((bank.transaction_date - cash.transaction_date).days)
    similarity = SequenceMatcher(None, bank_ref, cash_ref).ratio()
    blank_reference = _is_blank_reference(bank.reference) or _is_blank_reference(
        cash.reference
    )

    if same_reference and same_date:
        return (1, -1.0, 0), "Exact", "Reference, date and amount agree", 1.0
    if same_reference:
        confidence = max(0.90, 0.98 - min(date_difference, 30) * 0.002)
        return (
            (2, -1.0, date_difference),
            "Probable",
            f"Reference and amount agree; dates differ by {date_difference} day(s)",
            confidence,
        )
    if same_date and similarity >= 0.75:
        return (
            (3, -similarity, 0),
            "Probable",
            "Amount and date agree; reference differs",
            min(0.99, round(similarity, 3)),
        )
    if date_difference <= 3 and similarity >= 0.85:
        confidence = max(0.85, similarity - date_difference * 0.02)
        return (
            (4, -similarity, date_difference),
            "Probable",
            f"Amount agrees; similar reference and {date_difference}-day date difference",
            round(confidence, 3),
        )
    if same_date and blank_reference:
        return (
            (5, 0.0, 0),
            "Probable",
            "Amount and date agree; one reference is blank or a placeholder",
            0.88,
        )
    return None


def _match_receipts(
    bank_receipts: list[Receipt], cashbook_receipts: list[Receipt]
) -> tuple[list[Match], list[Receipt], list[Receipt]]:
    available_cash = set(range(len(cashbook_receipts)))
    matches: list[Match] = []
    matched_bank_rows: set[int] = set()

    for bank_index, bank in enumerate(bank_receipts):
        candidates = []
        for cash_index in sorted(available_cash):
            evidence = _candidate_evidence(bank, cashbook_receipts[cash_index])
            if evidence is not None:
                key, status, reason, confidence = evidence
                candidates.append(
                    (key, cash_index, status, reason, confidence)
                )
        if not candidates:
            continue

        candidates.sort(key=lambda item: item[0])
        best = candidates[0]
        tied = [item for item in candidates if item[0] == best[0]]
        if len(tied) > 1:
            # Equal-evidence candidates must remain unresolved for manual review.
            continue

        _, cash_index, status, reason, confidence = best
        matches.append(
            Match(
                bank=bank,
                cashbook=cashbook_receipts[cash_index],
                status=status,
                reason=reason,
                confidence=confidence,
            )
        )
        matched_bank_rows.add(bank_index)
        available_cash.remove(cash_index)

    matches.sort(
        key=lambda item: (
            item.bank.month,
            item.bank.transaction_date,
            item.bank.pdf_row,
        )
    )
    for index, match in enumerate(matches, start=1):
        match.match_id = f"M{index:04d}"

    bank_outstanding = [
        receipt
        for index, receipt in enumerate(bank_receipts)
        if index not in matched_bank_rows
    ]
    cashbook_outstanding = [
        receipt
        for index, receipt in enumerate(cashbook_receipts)
        if index in available_cash
    ]
    sort_key = lambda item: (item.month, item.transaction_date, item.pdf_row)
    bank_outstanding.sort(key=sort_key)
    cashbook_outstanding.sort(key=sort_key)
    return matches, bank_outstanding, cashbook_outstanding


def _set_title(ws, title: str, last_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=_NAVY)
    cell.font = Font(name="Carlito", color=_WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 27


def _style_header(ws, row: int, start_column: int, end_column: int) -> None:
    for column in range(start_column, end_column + 1):
        cell = ws.cell(row, column)
        cell.fill = PatternFill("solid", fgColor=_BLUE)
        cell.font = Font(name="Carlito", size=11, color=_WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER


def _style_section_bar(ws, row: int, start_column: int, end_column: int) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=start_column,
        end_row=row,
        end_column=end_column,
    )
    cell = ws.cell(row, start_column)
    cell.fill = PatternFill("solid", fgColor=_BLUE)
    cell.font = Font(name="Carlito", size=11, color=_WHITE, bold=True)
    cell.alignment = Alignment(horizontal="left")


def _style_body(ws, start_row: int, end_row: int, start_column: int, end_column: int):
    if end_row < start_row:
        return
    for row in ws.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        for cell in row:
            cell.font = Font(name="Carlito", size=11)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_table(ws, name: str, header_row: int, last_row: int, last_column: int) -> None:
    if last_row <= header_row:
        return
    from openpyxl.utils import get_column_letter

    ref = (
        f"A{header_row}:{get_column_letter(last_column)}{last_row}"
    )
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _write_matched_sheet(ws, matches: list[Match]) -> int:
    headers = [
        "Month",
        "Match ID",
        "Bank Reference",
        "Bank Date",
        "Bank Amount",
        "Bank PDF Row",
        "Bank Source Page",
        "Cash-book Reference",
        "Cash-book Date",
        "Cash-book Amount",
        "Cash-book PDF Row",
        "Cash-book Source Page",
        "Match Status",
        "Amount Difference",
        "Matching Reason",
        "Confidence",
        "Action",
        "Reviewer Decision",
    ]
    _set_title(ws, "MATCHED RECEIPTS", len(headers))
    ws.cell(
        2,
        1,
        "Exact matches can reconcile. Probable matches remain here for reviewer confirmation.",
    )
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=_SUBTITLE)
    ws.cell(2, 1).font = Font(name="Carlito", size=11)
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)
    _style_header(ws, 4, 1, len(headers))

    for row_number, match in enumerate(matches, start=5):
        values = [
            match.bank.month,
            match.match_id,
            match.bank.reference,
            match.bank.transaction_date,
            float(match.bank.amount),
            match.bank.pdf_row,
            match.bank.source_page,
            match.cashbook.reference,
            match.cashbook.transaction_date,
            float(match.cashbook.amount),
            match.cashbook.pdf_row,
            match.cashbook.source_page,
            match.status,
            f"=ROUND(E{row_number}-J{row_number},2)",
            match.reason,
            match.confidence,
            "Can reconcile"
            if match.status == "Exact"
            else "Review before reconciling",
            "",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
        fill = PatternFill(
            "solid", fgColor=_GREEN if match.status == "Exact" else _AMBER
        )
        for column in range(1, len(headers) + 1):
            ws.cell(row_number, column).fill = fill

    last_row = 4 + len(matches)
    _style_body(ws, 5, last_row, 1, len(headers))
    for row in range(5, last_row + 1):
        ws.cell(row, 1).number_format = _MONTH_FORMAT
        ws.cell(row, 4).number_format = _DATE_FORMAT
        ws.cell(row, 5).number_format = _CURRENCY_FORMAT
        ws.cell(row, 9).number_format = _DATE_FORMAT
        ws.cell(row, 10).number_format = _CURRENCY_FORMAT
        ws.cell(row, 14).number_format = _CURRENCY_FORMAT
        ws.cell(row, 16).number_format = "0.0%"
    widths = [13, 12, 22, 15, 17, 13, 16, 22, 15, 17, 16, 18, 14, 17, 42, 12, 25, 20]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index) if index <= 26 else "A"].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:R{max(4, last_row)}"
    _add_table(ws, "MatchedReceipts", 4, last_row, len(headers))
    ws.sheet_view.showGridLines = False
    return max(5, last_row)


def _write_bank_outstanding_sheet(ws, receipts: list[Receipt]) -> int:
    headers = ["Month", "Date", "Reference", "Amount", "PDF Row", "Source Page", "Action"]
    _set_title(ws, "BANK RECEIPTS OUTSTANDING", len(headers))
    ws.cell(2, 1, "Unmatched Section 2 receipts, arranged by month and date ascending.")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=_SUBTITLE)
    ws.cell(2, 1).font = Font(name="Carlito", size=11)
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)
    _style_header(ws, 4, 1, len(headers))

    for row_number, receipt in enumerate(receipts, start=5):
        values = [
            receipt.month,
            receipt.transaction_date,
            receipt.reference,
            float(receipt.amount),
            receipt.pdf_row,
            receipt.source_page,
            "Record in Cash Book / Investigate",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)

    last_row = 4 + len(receipts)
    _style_body(ws, 5, last_row, 1, len(headers))
    for row in range(5, last_row + 1):
        ws.cell(row, 1).number_format = _MONTH_FORMAT
        ws.cell(row, 2).number_format = _DATE_FORMAT
        ws.cell(row, 4).number_format = _CURRENCY_FORMAT
    for column, width in zip("ABCDEFG", [13, 15, 24, 17, 12, 16, 34]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:G{max(4, last_row)}"
    _add_table(ws, "BankOutstandingReceipts", 4, last_row, len(headers))
    total_row = last_row + 2
    ws.cell(total_row, 1, "Total Outstanding")
    ws.cell(total_row, 4, f"=SUM(D5:D{max(5, last_row)})")
    _style_body(ws, total_row, total_row, 1, len(headers))
    for column in range(1, len(headers) + 1):
        ws.cell(total_row, column).fill = PatternFill("solid", fgColor=_PALE_BLUE)
        ws.cell(total_row, column).font = Font(name="Carlito", size=11, bold=True)
    ws.cell(total_row, 4).number_format = _CURRENCY_FORMAT
    ws.sheet_view.showGridLines = False
    return max(5, last_row)


def _write_cashbook_outstanding_sheet(ws, receipts: list[Receipt]) -> int:
    headers = [
        "Month",
        "Date",
        "Reference",
        "Original Amount",
        "Reversal Amount",
        "PDF Row",
        "Source Page",
        "Action",
    ]
    _set_title(ws, "CASH-BOOK RECEIPTS OUTSTANDING", len(headers))
    ws.cell(2, 1, "Unmatched Section 4 receipts, arranged by month and date ascending.")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=_SUBTITLE)
    ws.cell(2, 1).font = Font(name="Carlito", size=11)
    for column, header in enumerate(headers, start=1):
        ws.cell(4, column, header)
    _style_header(ws, 4, 1, len(headers))

    for row_number, receipt in enumerate(receipts, start=5):
        values = [
            receipt.month,
            receipt.transaction_date,
            receipt.reference,
            float(receipt.amount),
            f"=-D{row_number}",
            receipt.pdf_row,
            receipt.source_page,
            "Reverse in Cash Book",
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)

    last_row = 4 + len(receipts)
    _style_body(ws, 5, last_row, 1, len(headers))
    for row in range(5, last_row + 1):
        ws.cell(row, 1).number_format = _MONTH_FORMAT
        ws.cell(row, 2).number_format = _DATE_FORMAT
        ws.cell(row, 4).number_format = _CURRENCY_FORMAT
        ws.cell(row, 5).number_format = _CURRENCY_FORMAT
    for column, width in zip("ABCDEFGH", [13, 15, 24, 18, 18, 12, 16, 25]):
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:H{max(4, last_row)}"
    _add_table(ws, "CashbookOutstandingReceipts", 4, last_row, len(headers))
    total_row = last_row + 2
    ws.cell(total_row, 1, "Total Outstanding")
    ws.cell(total_row, 4, f"=SUM(D5:D{max(5, last_row)})")
    ws.cell(total_row, 5, f"=SUM(E5:E{max(5, last_row)})")
    _style_body(ws, total_row, total_row, 1, len(headers))
    for column in range(1, len(headers) + 1):
        ws.cell(total_row, column).fill = PatternFill("solid", fgColor=_PALE_BLUE)
        ws.cell(total_row, column).font = Font(name="Carlito", size=11, bold=True)
    ws.cell(total_row, 4).number_format = _CURRENCY_FORMAT
    ws.cell(total_row, 5).number_format = _CURRENCY_FORMAT
    ws.sheet_view.showGridLines = False
    return max(5, last_row)


def _write_summary_sheet(
    ws,
    metadata: _Metadata,
    bank_receipts: list[Receipt],
    cashbook_receipts: list[Receipt],
    matches: list[Match],
    bank_outstanding: list[Receipt],
    cashbook_outstanding: list[Receipt],
    matched_end: int,
    bank_end: int,
    cash_end: int,
) -> None:
    _set_title(ws, "RECEIPT RECONCILIATION - F.O. 30", 10)
    ws["A3"] = "Source PDF"
    ws["B3"] = metadata.source_name
    ws.merge_cells("B3:D3")
    ws["A4"] = "Bank"
    ws["B4"] = metadata.bank
    ws.merge_cells("B4:D4")
    ws["A5"] = "Branch"
    ws["B5"] = metadata.branch
    ws.merge_cells("B5:D5")
    ws["F3"] = "Account Number"
    ws["G3"] = metadata.account_number
    ws.merge_cells("G3:J3")
    ws["F4"] = "Period"
    ws["G4"] = f"{metadata.period_from:%d-%b-%Y} to {metadata.period_to:%d-%b-%Y}"
    ws.merge_cells("G4:J4")
    ws["F5"] = "Scope"
    ws["G5"] = "Receipts only - Sections 2 and 4"
    ws.merge_cells("G5:J5")
    for cell_ref in ("A3", "A4", "A5", "F3", "F4", "F5"):
        ws[cell_ref].fill = PatternFill("solid", fgColor=_PALE_BLUE)
        ws[cell_ref].font = Font(name="Carlito", size=11, bold=True, color=_NAVY)

    _style_section_bar(ws, 8, 1, 5)
    ws["A8"] = "RECEIPT POPULATION"
    population_headers = ["Population", "Count", "Extracted Total", "Printed Total", "Difference"]
    for column, header in enumerate(population_headers, start=1):
        ws.cell(9, column, header)
    _style_header(ws, 9, 1, 5)
    ws["A10"] = "Section 2 - Bank Statement"
    ws["B10"] = (
        f"=COUNTA('Matched'!$B$5:$B${matched_end})+"
        f"COUNTA('Bank Outstanding'!$C$5:$C${bank_end})"
    )
    ws["C10"] = (
        f"=SUM('Matched'!$E$5:$E${matched_end})+"
        f"SUM('Bank Outstanding'!$D$5:$D${bank_end})"
    )
    ws["D10"] = float(metadata.printed_totals[2])
    ws["E10"] = "=ROUND(C10-D10,2)"
    ws["A11"] = "Section 4 - Cash Book"
    ws["B11"] = (
        f"=COUNTA('Matched'!$B$5:$B${matched_end})+"
        f"COUNTA('Cashbook Outstanding'!$C$5:$C${cash_end})"
    )
    ws["C11"] = (
        f"=SUM('Matched'!$J$5:$J${matched_end})+"
        f"SUM('Cashbook Outstanding'!$D$5:$D${cash_end})"
    )
    ws["D11"] = float(metadata.printed_totals[4])
    ws["E11"] = "=ROUND(C11-D11,2)"
    _style_body(ws, 10, 11, 1, 5)

    _style_section_bar(ws, 13, 1, 5)
    ws["A13"] = "MATCHING RESULTS"
    match_headers = ["Classification", "Count", "Bank Amount", "Cash-book Amount", "Difference"]
    for column, header in enumerate(match_headers, start=1):
        ws.cell(14, column, header)
    _style_header(ws, 14, 1, 5)
    match_rows = [
        (
            "Exact matches",
            f'=COUNTIF(\'Matched\'!$M$5:$M${matched_end},"Exact")',
            f'=SUMIF(\'Matched\'!$M$5:$M${matched_end},"Exact",\'Matched\'!$E$5:$E${matched_end})',
            f'=SUMIF(\'Matched\'!$M$5:$M${matched_end},"Exact",\'Matched\'!$J$5:$J${matched_end})',
            "=ROUND(C15-D15,2)",
        ),
        (
            "Probable matches",
            f'=COUNTIF(\'Matched\'!$M$5:$M${matched_end},"Probable")',
            f'=SUMIF(\'Matched\'!$M$5:$M${matched_end},"Probable",\'Matched\'!$E$5:$E${matched_end})',
            f'=SUMIF(\'Matched\'!$M$5:$M${matched_end},"Probable",\'Matched\'!$J$5:$J${matched_end})',
            "=ROUND(C16-D16,2)",
        ),
        (
            "Bank outstanding",
            f"=COUNTA('Bank Outstanding'!$C$5:$C${bank_end})",
            f"=SUM('Bank Outstanding'!$D$5:$D${bank_end})",
            0,
            "=ROUND(C17-D17,2)",
        ),
        (
            "Cash-book outstanding",
            f"=COUNTA('Cashbook Outstanding'!$C$5:$C${cash_end})",
            0,
            f"=SUM('Cashbook Outstanding'!$D$5:$D${cash_end})",
            "=ROUND(C18-D18,2)",
        ),
        (
            "Negative cash-book reversals",
            f"=COUNTA('Cashbook Outstanding'!$C$5:$C${cash_end})",
            0,
            f"=SUM('Cashbook Outstanding'!$E$5:$E${cash_end})",
            f"=SUM('Cashbook Outstanding'!$D$5:$D${cash_end})+"
            f"SUM('Cashbook Outstanding'!$E$5:$E${cash_end})",
        ),
    ]
    for row_number, values in enumerate(match_rows, start=15):
        for column, value in enumerate(values, start=1):
            ws.cell(row_number, column, value)
    _style_body(ws, 15, 19, 1, 5)
    for column in range(1, 6):
        ws.cell(15, column).fill = PatternFill("solid", fgColor=_GREEN)
        ws.cell(16, column).fill = PatternFill("solid", fgColor=_AMBER)

    _style_section_bar(ws, 8, 7, 10)
    ws["G8"] = "ZERO-DIFFERENCE QUALITY CONTROLS"
    qc_headers = ["Control", "Expected", "Actual", "Result"]
    for column, header in enumerate(qc_headers, start=7):
        ws.cell(9, column, header)
    _style_header(ws, 9, 7, 10)
    controls = [
        ("Section 2 extraction", 0, "=E10"),
        ("Section 4 extraction", 0, "=E11"),
        ("Section 2 allocation", 0, "=ROUND(C15+C16+C17-D10,2)"),
        ("Section 4 allocation", 0, "=ROUND(D15+D16+D18-D11,2)"),
        (
            "Receipt count allocation",
            0,
            "=2*(B15+B16)+B17+B18-(B10+B11)",
        ),
        (
            "Matched amount differences",
            0,
            f"=SUMPRODUCT(ABS('Matched'!$N$5:$N${matched_end}))",
        ),
        ("Reversal absolute total", 0, "=ROUND(ABS(D19)-D18,2)"),
        ("Reversal net total", 0, "=ROUND(D18+D19,2)"),
    ]
    for row_number, (label, expected, actual) in enumerate(controls, start=10):
        ws.cell(row_number, 7, label)
        ws.cell(row_number, 8, expected)
        ws.cell(row_number, 9, actual)
        ws.cell(
            row_number,
            10,
            f'=IF(ABS(I{row_number}-H{row_number})<0.01,"PASS","CHECK")',
        )
    _style_body(ws, 10, 17, 7, 10)
    green_fill = PatternFill("solid", fgColor=_GREEN)
    red_fill = PatternFill("solid", fgColor=_RED)
    ws.conditional_formatting.add(
        "J10:J17",
        CellIsRule(operator="equal", formula=['"PASS"'], fill=green_fill),
    )
    ws.conditional_formatting.add(
        "J10:J17",
        CellIsRule(operator="equal", formula=['"CHECK"'], fill=red_fill),
    )

    months = sorted({receipt.month for receipt in bank_receipts + cashbook_receipts})
    monthly_title_row = 22
    monthly_header_row = 23
    _style_section_bar(ws, monthly_title_row, 1, 5)
    ws.cell(monthly_title_row, 1, "MONTHLY RECEIPT ANALYSIS - ASCENDING")
    monthly_headers = [
        "Month",
        "Bank Receipt Count",
        "Bank Receipt Total",
        "Cash-book Receipt Count",
        "Cash-book Receipt Total",
    ]
    for column, header in enumerate(monthly_headers, start=1):
        ws.cell(monthly_header_row, column, header)
    _style_header(ws, monthly_header_row, 1, 5)

    for row_number, month in enumerate(months, start=24):
        ws.cell(row_number, 1, month)
        ws.cell(
            row_number,
            2,
            f'=COUNTIFS(\'Matched\'!$D$5:$D${matched_end},">="&$A{row_number},'
            f'\'Matched\'!$D$5:$D${matched_end},"<"&EDATE($A{row_number},1))+'
            f'=COUNTIFS(\'Bank Outstanding\'!$B$5:$B${bank_end},">="&$A{row_number},'
            f'\'Bank Outstanding\'!$B$5:$B${bank_end},"<"&EDATE($A{row_number},1))',
        )
        # Remove the second leading "=" after joining two complete COUNTIFS formulas.
        ws.cell(row_number, 2).value = ws.cell(row_number, 2).value.replace("+=", "+")
        ws.cell(
            row_number,
            3,
            f'=SUMIFS(\'Matched\'!$E$5:$E${matched_end},'
            f'\'Matched\'!$D$5:$D${matched_end},">="&$A{row_number},'
            f'\'Matched\'!$D$5:$D${matched_end},"<"&EDATE($A{row_number},1))+'
            f'SUMIFS(\'Bank Outstanding\'!$D$5:$D${bank_end},'
            f'\'Bank Outstanding\'!$B$5:$B${bank_end},">="&$A{row_number},'
            f'\'Bank Outstanding\'!$B$5:$B${bank_end},"<"&EDATE($A{row_number},1))',
        )
        ws.cell(
            row_number,
            4,
            f'=COUNTIFS(\'Matched\'!$I$5:$I${matched_end},">="&$A{row_number},'
            f'\'Matched\'!$I$5:$I${matched_end},"<"&EDATE($A{row_number},1))+'
            f'COUNTIFS(\'Cashbook Outstanding\'!$B$5:$B${cash_end},">="&$A{row_number},'
            f'\'Cashbook Outstanding\'!$B$5:$B${cash_end},"<"&EDATE($A{row_number},1))',
        )
        ws.cell(
            row_number,
            5,
            f'=SUMIFS(\'Matched\'!$J$5:$J${matched_end},'
            f'\'Matched\'!$I$5:$I${matched_end},">="&$A{row_number},'
            f'\'Matched\'!$I$5:$I${matched_end},"<"&EDATE($A{row_number},1))+'
            f'SUMIFS(\'Cashbook Outstanding\'!$D$5:$D${cash_end},'
            f'\'Cashbook Outstanding\'!$B$5:$B${cash_end},">="&$A{row_number},'
            f'\'Cashbook Outstanding\'!$B$5:$B${cash_end},"<"&EDATE($A{row_number},1))',
        )

    monthly_last = 23 + len(months)
    total_row = monthly_last + 2
    ws.cell(total_row, 1, "Total")
    for column in range(2, 6):
        letter = chr(64 + column)
        ws.cell(column=column, row=total_row, value=f"=SUM({letter}24:{letter}{monthly_last})")
    _style_body(ws, 24, monthly_last, 1, 5)
    _style_body(ws, total_row, total_row, 1, 5)
    for cell in ws[total_row]:
        if cell.column <= 5:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=_PALE_BLUE)

    for row in range(10, 20):
        for column in (3, 4, 5):
            ws.cell(row, column).number_format = _CURRENCY_FORMAT
    for row in range(10, 18):
        ws.cell(row, 8).number_format = _CURRENCY_FORMAT
        ws.cell(row, 9).number_format = _CURRENCY_FORMAT
    for row in range(24, monthly_last + 1):
        ws.cell(row, 1).number_format = _MONTH_FORMAT
        ws.cell(row, 3).number_format = _CURRENCY_FORMAT
        ws.cell(row, 5).number_format = _CURRENCY_FORMAT
    ws.cell(total_row, 3).number_format = _CURRENCY_FORMAT
    ws.cell(total_row, 5).number_format = _CURRENCY_FORMAT

    for column, width in {
        "A": 31,
        "B": 19,
        "C": 20,
        "D": 20,
        "E": 18,
        "F": 3,
        "G": 30,
        "H": 14,
        "I": 18,
        "J": 14,
    }.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A8"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:9"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _validate_result(
    metadata: _Metadata,
    bank_receipts: list[Receipt],
    cashbook_receipts: list[Receipt],
    matches: list[Match],
    bank_outstanding: list[Receipt],
    cashbook_outstanding: list[Receipt],
) -> None:
    bank_total = sum((receipt.amount for receipt in bank_receipts), Decimal("0.00"))
    cash_total = sum((receipt.amount for receipt in cashbook_receipts), Decimal("0.00"))
    matched_bank_total = sum((match.bank.amount for match in matches), Decimal("0.00"))
    matched_cash_total = sum(
        (match.cashbook.amount for match in matches), Decimal("0.00")
    )
    bank_outstanding_total = sum(
        (receipt.amount for receipt in bank_outstanding), Decimal("0.00")
    )
    cash_outstanding_total = sum(
        (receipt.amount for receipt in cashbook_outstanding), Decimal("0.00")
    )

    checks = [
        (bank_total == metadata.printed_totals[2], "Section 2 printed total"),
        (cash_total == metadata.printed_totals[4], "Section 4 printed total"),
        (
            matched_bank_total + bank_outstanding_total == bank_total,
            "Section 2 allocation",
        ),
        (
            matched_cash_total + cash_outstanding_total == cash_total,
            "Section 4 allocation",
        ),
        (
            2 * len(matches) + len(bank_outstanding) + len(cashbook_outstanding)
            == len(bank_receipts) + len(cashbook_receipts),
            "receipt count allocation",
        ),
        (
            all(match.bank.amount == match.cashbook.amount for match in matches),
            "matched amount differences",
        ),
        (
            all(
                bank_outstanding[index - 1].month <= bank_outstanding[index].month
                for index in range(1, len(bank_outstanding))
            ),
            "bank outstanding sort order",
        ),
        (
            all(
                cashbook_outstanding[index - 1].month
                <= cashbook_outstanding[index].month
                for index in range(1, len(cashbook_outstanding))
            ),
            "cash-book outstanding sort order",
        ),
    ]
    failed = [name for passed, name in checks if not passed]
    if failed:
        raise ReceiptReconciliationError(
            "Receipt reconciliation quality control failed: " + ", ".join(failed)
        )


def _build_workbook(
    metadata: _Metadata,
    bank_receipts: list[Receipt],
    cashbook_receipts: list[Receipt],
    matches: list[Match],
    bank_outstanding: list[Receipt],
    cashbook_outstanding: list[Receipt],
) -> Workbook:
    workbook = Workbook()
    summary_ws = workbook.active
    summary_ws.title = "Summary"
    matched_ws = workbook.create_sheet("Matched")
    bank_ws = workbook.create_sheet("Bank Outstanding")
    cashbook_ws = workbook.create_sheet("Cashbook Outstanding")

    matched_end = _write_matched_sheet(matched_ws, matches)
    bank_end = _write_bank_outstanding_sheet(bank_ws, bank_outstanding)
    cash_end = _write_cashbook_outstanding_sheet(cashbook_ws, cashbook_outstanding)
    _write_summary_sheet(
        summary_ws,
        metadata,
        bank_receipts,
        cashbook_receipts,
        matches,
        bank_outstanding,
        cashbook_outstanding,
        matched_end,
        bank_end,
        cash_end,
    )
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    return workbook


def _validate_saved_workbook(path: str) -> None:
    workbook = load_workbook(path, data_only=False, read_only=False)
    try:
        if workbook.sheetnames != REQUIRED_SHEETS:
            raise ReceiptReconciliationError(
                "Generated workbook does not have the required four sheets in order."
            )
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        upper = cell.value.upper()
                        if any(error in upper for error in ("#REF!", "#DIV/0!", "#VALUE!")):
                            raise ReceiptReconciliationError(
                                f"Formula error in {worksheet.title}!{cell.coordinate}."
                            )
    finally:
        workbook.close()


def default_output_path(pdf_path: str, account_number: str | None = None) -> str:
    folder = os.path.dirname(os.path.abspath(pdf_path))
    account = re.sub(r"[^A-Z0-9-]+", "", str(account_number or ""), flags=re.IGNORECASE)
    suffix = f"_Account_{account}" if account else ""
    return os.path.join(folder, f"Receipt_Reconciliation{suffix}.xlsx")


def inspect_receipt_pdf(pdf_path: str) -> tuple[str, str]:
    """Return account number and the default output path after full PDF validation."""
    metadata, _, _ = _extract_pdf(pdf_path)
    return metadata.account_number, default_output_path(pdf_path, metadata.account_number)


def create_receipt_reconciliation(pdf_path: str, output_path: str) -> ReconciliationResult:
    """Create the four-sheet receipt reconciliation workbook."""
    pdf_path = os.path.abspath(pdf_path)
    output_path = os.path.abspath(output_path)
    if not os.path.isfile(pdf_path):
        raise ReceiptReconciliationError("Select an existing F.O. 30 PDF.")
    if Path(pdf_path).suffix.lower() != ".pdf":
        raise ReceiptReconciliationError("The receipt reconciliation source must be a PDF.")
    if Path(output_path).suffix.lower() != ".xlsx":
        raise ReceiptReconciliationError("The output file must use the .xlsx extension.")
    if os.path.normcase(pdf_path) == os.path.normcase(output_path):
        raise ReceiptReconciliationError("The output path cannot replace the source PDF.")

    metadata, bank_receipts, cashbook_receipts = _extract_pdf(pdf_path)
    matches, bank_outstanding, cashbook_outstanding = _match_receipts(
        bank_receipts, cashbook_receipts
    )
    _validate_result(
        metadata,
        bank_receipts,
        cashbook_receipts,
        matches,
        bank_outstanding,
        cashbook_outstanding,
    )
    workbook = _build_workbook(
        metadata,
        bank_receipts,
        cashbook_receipts,
        matches,
        bank_outstanding,
        cashbook_outstanding,
    )

    output_folder = os.path.dirname(output_path)
    os.makedirs(output_folder, exist_ok=True)
    temp_path = ""
    try:
        file_descriptor, temp_path = tempfile.mkstemp(
            prefix=".receipt_recon_",
            suffix=".xlsx",
            dir=output_folder,
        )
        os.close(file_descriptor)
        workbook.save(temp_path)
        _validate_saved_workbook(temp_path)
        os.replace(temp_path, output_path)
        temp_path = ""
    finally:
        workbook.close()
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except OSError:
                pass

    return ReconciliationResult(
        output_path=output_path,
        bank_count=len(bank_receipts),
        cashbook_count=len(cashbook_receipts),
        exact_count=sum(match.status == "Exact" for match in matches),
        probable_count=sum(match.status == "Probable" for match in matches),
        bank_outstanding_count=len(bank_outstanding),
        cashbook_outstanding_count=len(cashbook_outstanding),
        bank_total=metadata.printed_totals[2],
        cashbook_total=metadata.printed_totals[4],
    )
